"""Exportul raportului lunar în Excel și PDF, pentru contabilă."""

import io
import sqlite3
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .calc import month_bounds, month_summary, payroll
from .db import get_settings

FONTS_DIR = Path(__file__).resolve().parent / "fonts"
MONEY = '#,##0.00 "lei"'
HEADER_FILL = PatternFill("solid", fgColor="0F172A")
ACCENT = "0EA5E9"


def summary_rows(summary: dict) -> list[tuple[str, float]]:
    taxes = summary["taxes"]
    return [
        ("Venituri din vânzări", summary["revenue"]),
        ("Cheltuieli operaționale", summary["operating_expenses"]),
        ("Salarii brute", summary["payroll"]["gross"]),
        ("CAM (plătit de firmă)", taxes["cam"]),
        ("Impozit microîntreprindere", taxes["micro_tax"]),
        ("Total costuri", summary["total_costs"]),
        ("Profit net", summary["profit"]),
        ("Impozit pe dividende (estimat)", summary["dividend_tax"]),
        ("Rămâne proprietarului", summary["owner_income"]),
    ]


def tax_rows(summary: dict) -> list[tuple[str, float]]:
    taxes = summary["taxes"]
    return [
        ("CAS — pensie (din salarii)", taxes["cas"]),
        ("CASS — sănătate (din salarii)", taxes["cass"]),
        ("Impozit pe venit (din salarii)", taxes["income_tax"]),
        ("CAM — contribuție firmă", taxes["cam"]),
        ("Impozit microîntreprindere", taxes["micro_tax"]),
        ("Total taxe", taxes["total"]),
    ]


def _write_table(sheet, headers: list[str], rows: list[list], money_cols: set[int]) -> None:
    sheet.append(headers)
    for cell in sheet[sheet.max_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        sheet.append(row)
        for index in money_cols:
            sheet.cell(row=sheet.max_row, column=index).number_format = MONEY
    for index, header in enumerate(headers, start=1):
        values = [len(str(header))] + [len(str(row[index - 1] or "")) for row in rows]
        sheet.column_dimensions[get_column_letter(index)].width = min(max(values) + 4, 48)
    sheet.freeze_panes = "A2"


def excel_report(conn: sqlite3.Connection, ym: str) -> bytes:
    settings = get_settings(conn)
    summary = month_summary(conn, ym, settings)
    start, end = month_bounds(ym)
    book = Workbook()

    sheet = book.active
    sheet.title = "Sumar"
    sheet["A1"] = f"{settings['company_name']} — Raport {summary['label']}"
    sheet["A1"].font = Font(bold=True, size=16, color=ACCENT)
    sheet.append([])
    _write_table(sheet, ["Indicator", "Valoare"],
                 [list(row) for row in summary_rows(summary)], {2})
    sheet.append([])
    _write_table(sheet, ["Taxe și contribuții", "Valoare"],
                 [list(row) for row in tax_rows(summary)], {2})
    sheet.freeze_panes = None
    sheet.column_dimensions["A"].width = 38
    sheet.column_dimensions["B"].width = 20

    sales = conn.execute(
        "SELECT date, description, product, quantity, amount FROM sales "
        "WHERE date BETWEEN ? AND ? ORDER BY date, id", (start, end)
    ).fetchall()
    _write_table(book.create_sheet("Vânzări"),
                 ["Data", "Descriere", "Produs", "Cantitate", "Sumă"],
                 [list(row) for row in sales], {5})

    expenses = conn.execute(
        "SELECT date, category, description, amount FROM expenses "
        "WHERE date BETWEEN ? AND ? ORDER BY date, id", (start, end)
    ).fetchall()
    _write_table(book.create_sheet("Cheltuieli"),
                 ["Data", "Categorie", "Descriere", "Sumă"],
                 [list(row) for row in expenses], {4})

    pay = payroll(conn, ym, settings)
    rows = [[line["name"], line["role"], line["gross"], line["cas"], line["cass"],
             line["income_tax"], line["net"], line["cam"], line["employer_cost"]]
            for line in pay["lines"]]
    totals = pay["totals"]
    rows.append(["TOTAL", "", totals["gross"], totals["cas"], totals["cass"],
                 totals["income_tax"], totals["net"], totals["cam"], totals["employer_cost"]])
    salaries = book.create_sheet("Salarii")
    _write_table(salaries, ["Angajat", "Funcție", "Brut", "CAS", "CASS", "Impozit",
                            "Net", "CAM", "Cost firmă"], rows, set(range(3, 10)))
    for cell in salaries[salaries.max_row]:
        cell.font = Font(bold=True)

    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _register_fonts() -> tuple[str, str]:
    try:
        pdfmetrics.registerFont(TTFont("DejaVu", FONTS_DIR / "DejaVuSans.ttf"))
        pdfmetrics.registerFont(TTFont("DejaVu-Bold", FONTS_DIR / "DejaVuSans-Bold.ttf"))
        return "DejaVu", "DejaVu-Bold"
    except Exception:
        return "Helvetica", "Helvetica-Bold"


def lei(value: float) -> str:
    return f"{value:,.2f}".replace(",", " ").replace(".", ",").replace(" ", ".") + " lei"


def pdf_report(conn: sqlite3.Connection, ym: str) -> bytes:
    regular, bold = _register_fonts()
    settings = get_settings(conn)
    summary = month_summary(conn, ym, settings)
    pay = payroll(conn, ym, settings)

    title = ParagraphStyle("title", fontName=bold, fontSize=20, leading=24,
                           textColor=colors.HexColor("#0f172a"))
    subtitle = ParagraphStyle("subtitle", fontName=regular, fontSize=11,
                              textColor=colors.HexColor("#0ea5e9"), spaceAfter=10)
    heading = ParagraphStyle("heading", fontName=bold, fontSize=13, spaceBefore=12,
                             spaceAfter=6, textColor=colors.HexColor("#0f172a"))

    def table(headers: list[str], rows: list[list], widths: list[float],
              total_row: bool = False) -> Table:
        result = Table([headers] + rows, colWidths=widths, repeatRows=1)
        style = [
            ("FONTNAME", (0, 0), (-1, -1), regular),
            ("FONTNAME", (0, 0), (-1, 0), bold),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f172a")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f1f5f9")]),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#cbd5e1")),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]
        if total_row:
            style += [("FONTNAME", (0, -1), (-1, -1), bold),
                      ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#e0f2fe"))]
        result.setStyle(TableStyle(style))
        return result

    story = [
        Paragraph(settings["company_name"], title),
        Paragraph(f"Raport financiar lunar — {summary['label']}", subtitle),
        Paragraph("Sumar", heading),
        table(["Indicator", "Valoare"],
              [[name, lei(value)] for name, value in summary_rows(summary)],
              [110 * mm, 60 * mm]),
        Paragraph("Taxe și contribuții", heading),
        table(["Taxă", "Valoare"],
              [[name, lei(value)] for name, value in tax_rows(summary)],
              [110 * mm, 60 * mm], total_row=True),
        Paragraph("Cheltuieli pe categorii", heading),
        table(["Categorie", "Valoare"],
              [[name, lei(value)] for name, value in summary["expenses_by_category"].items()]
              + [["Total", lei(summary["operating_expenses"])]],
              [110 * mm, 60 * mm], total_row=True),
        Paragraph("Salarii", heading),
        table(["Angajat", "Brut", "Rețineri", "Net", "Cost firmă"],
              [[line["name"], lei(line["gross"]),
                lei(line["cas"] + line["cass"] + line["income_tax"]),
                lei(line["net"]), lei(line["employer_cost"])] for line in pay["lines"]]
              + [["Total", lei(pay["totals"]["gross"]),
                  lei(pay["totals"]["cas"] + pay["totals"]["cass"] + pay["totals"]["income_tax"]),
                  lei(pay["totals"]["net"]), lei(pay["totals"]["employer_cost"])]],
              [50 * mm, 30 * mm, 30 * mm, 30 * mm, 30 * mm], total_row=True),
        Spacer(1, 10 * mm),
        Paragraph("Generat automat de aplicația internă Testify Hardware.",
                  ParagraphStyle("foot", fontName=regular, fontSize=8,
                                 textColor=colors.HexColor("#64748b"))),
    ]

    buffer = io.BytesIO()
    SimpleDocTemplate(buffer, pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm,
                      topMargin=18 * mm, bottomMargin=18 * mm,
                      title=f"Raport {summary['label']}").build(story)
    return buffer.getvalue()
