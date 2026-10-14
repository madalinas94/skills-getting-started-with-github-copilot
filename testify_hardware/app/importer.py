"""Importul datelor firmei din fișierele pe care le folosea până acum (CSV și raportul Excel).

Fiecare tabel este recunoscut după coloanele lui, nu după numele fișierului, așa că merg
atât exporturile CSV (produse, clienți, angajați, vânzări, cheltuieli, inventar), cât și
foile din raportul contabil (Vanzari, Achizitii, Cheltuieli, Salarii, Stocuri, Clienti).
"""

import csv
import io
import re
import sqlite3
import unicodedata
from datetime import date, datetime, timedelta

from openpyxl import load_workbook

from .calc import month_bounds
from .db import DEFAULT_SUPPLIER_TERMS

CATEGORY_MAP = {
    "chirie": "Chirie",
    "utilitati": "Utilități",
    "contabilitate": "Contabilitate",
    "software": "Software și abonamente",
    "marketing": "Marketing",
    "transport": "Transport",
    "consumabile": "Ambalaje și consumabile",
    "mentenanta": "Mentenanță și reparații",
    "banca": "Comisioane bancare și POS",
    "echipamente_it": "Echipamente IT și mobilier",
}
DATA_TABLES = ("sales", "expenses", "purchases", "salary_history", "employees", "products", "clients",
               "suppliers")


def norm(text) -> str:
    """„Preț unitar (fără TVA)” → „pret_unitar_fara_tva”."""
    text = unicodedata.normalize("NFKD", str(text or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "_", text.lower()).strip("_")


def to_date(value) -> str | None:
    if value in (None, ""):
        return None
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(text[:10], fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    raise ValueError(f"Dată necunoscută: {text}")


def to_float(value) -> float:
    if value in (None, ""):
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(" ", "")
    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".")
    elif "," in text:
        text = text.replace(",", ".")
    return float(text)


def is_number(value) -> bool:
    try:
        to_float(value)
        return value not in (None, "") and not str(value).startswith("=")
    except ValueError:
        return False


class Table:
    def __init__(self, name: str, headers: list, rows: list[list]):
        self.name = name
        self.headers = [norm(h) for h in headers]
        self.rows = [dict(zip(self.headers, row)) for row in rows
                     if any(cell not in (None, "") for cell in row)]

    def has(self, *columns) -> bool:
        return all(any(h == c or h.startswith(c) for h in self.headers) for c in columns)

    def col(self, prefix: str) -> str | None:
        return next((h for h in self.headers if h == prefix), None) or \
            next((h for h in self.headers if h.startswith(prefix)), None)

    def kind(self) -> str | None:
        if self.has("nr_factura", "cantitate"):
            return "sales"
        if self.has("nr_nir", "cod_produs"):
            return "purchases"
        if self.has("cod_produs", "pret_vanzare"):
            return "products"
        if self.has("cod_produs", "stoc_initial"):
            return "stock"
        if self.has("client", "tip_client"):
            return "clients"
        if self.has("categorie", "suma_fara_tva"):
            return "expenses"
        if self.has("nume", "salariu_brut"):
            return "payroll"
        return None


def tables_from_csv(filename: str, content: bytes) -> list[Table]:
    text = content.decode("utf-8-sig")
    dialect = csv.Sniffer().sniff(text[:2000], delimiters=",;\t")
    rows = list(csv.reader(io.StringIO(text), dialect))
    return [Table(filename, rows[0], rows[1:])] if rows else []


def tables_from_xlsx(filename: str, content: bytes) -> tuple[list[Table], str | None]:
    book = load_workbook(io.BytesIO(content), data_only=False)
    tables, reference = [], None
    for sheet in book.worksheets:
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        rows = [row for row in rows if any(cell not in (None, "") for cell in row)]
        if not rows:
            continue
        # Data de referință pentru încasări, din foaia de ipoteze a contabilului.
        for row in rows:
            if row and norm(row[0]).startswith("data_de_referinta") and len(row) > 1:
                reference = to_date(row[1])
        # Antetul este primul rând care arată ca un antet de tabel.
        for index, row in enumerate(rows[:5]):
            table = Table(f"{filename}:{sheet.title}", row, rows[index + 1:])
            if table.kind():
                tables.append(table)
                break
    return tables, reference


class Importer:
    ORDER = ["clients", "products", "stock", "purchases", "sales", "expenses", "payroll"]

    def __init__(self, conn: sqlite3.Connection, reference_date: str | None = None):
        self.conn = conn
        self.reference = reference_date
        self.counts: dict[str, int] = {}
        self.warnings: list[str] = []

    def add(self, key: str, n: int = 1) -> None:
        self.counts[key] = self.counts.get(key, 0) + n

    # ---------- ajutoare ----------

    def product_id(self, name: str, code: str | None = None, category: str = "",
                   cost: float = 0, price: float = 0) -> int:
        row = None
        if code:
            row = self.conn.execute("SELECT id FROM products WHERE code = ?", (code,)).fetchone()
        if not row and name:
            row = self.conn.execute("SELECT id FROM products WHERE name = ?", (name,)).fetchone()
        if row:
            return row["id"]
        if not code:
            count = self.conn.execute("SELECT COUNT(*) FROM products").fetchone()[0]
            code = f"P{1000 + count + 1}"
        cursor = self.conn.execute(
            "INSERT INTO products(code, name, category, cost, sale_price) VALUES (?, ?, ?, ?, ?)",
            (code, name, category, cost, price))
        self.add("products")
        return cursor.lastrowid

    def client(self, name: str) -> sqlite3.Row:
        name = (name or "Persoana fizica").strip()
        row = self.conn.execute("SELECT * FROM clients WHERE name = ?", (name,)).fetchone()
        if row:
            return row
        self.conn.execute("INSERT INTO clients(name) VALUES (?)", (name,))
        self.add("clients")
        return self.conn.execute("SELECT * FROM clients WHERE name = ?", (name,)).fetchone()

    def supplier_terms(self, name: str) -> int:
        """Termenul de plată al furnizorului; furnizorii noi primesc termenul implicit."""
        name = (name or "").strip()
        if not name:
            return 0
        row = self.conn.execute("SELECT payment_terms FROM suppliers WHERE name = ?", (name,)).fetchone()
        if row:
            return row["payment_terms"]
        self.conn.execute("INSERT INTO suppliers(name, payment_terms) VALUES (?, ?)",
                          (name, DEFAULT_SUPPLIER_TERMS))
        self.add("suppliers")
        return DEFAULT_SUPPLIER_TERMS

    def category(self, raw: str) -> str:
        key = norm(raw)
        if key in CATEGORY_MAP:
            return CATEGORY_MAP[key]
        for row in self.conn.execute("SELECT name FROM categories"):
            if norm(row["name"]) == key:
                return row["name"]
        name = str(raw).strip() or "Altele"
        self.conn.execute("INSERT OR IGNORE INTO categories(name, color) VALUES (?, '#94a3b8')", (name,))
        return name

    # ---------- tabele ----------

    def clients(self, table: Table) -> None:
        terms_col = table.col("termen_plata")
        for row in table.rows:
            name = str(row.get("client") or "").strip()
            if not name or norm(name) == "total":
                continue
            values = (str(row.get("tip_client") or "B2C"), str(row.get("oras") or "").strip("- "),
                      str(row.get("agent_responsabil") or "").strip("- "),
                      int(to_float(row.get(terms_col))) if terms_col else 0)
            exists = self.conn.execute("SELECT 1 FROM clients WHERE name = ?", (name,)).fetchone()
            self.conn.execute(
                "INSERT INTO clients(name, type, city, agent, payment_terms) VALUES (?, ?, ?, ?, ?) "
                "ON CONFLICT(name) DO UPDATE SET type = excluded.type, city = excluded.city, "
                "agent = excluded.agent, payment_terms = excluded.payment_terms", (name, *values))
            if not exists:
                self.add("clients")

    def products(self, table: Table) -> None:
        for row in table.rows:
            code = str(row.get("cod_produs") or "").strip()
            if not code:
                continue
            values = (str(row.get("produs") or code), str(row.get("categorie") or ""),
                      str(row.get("furnizor") or ""), to_float(row.get(table.col("pret_vanzare"))),
                      to_float(row.get(table.col("cost_achizitie"))),
                      int(to_float(row.get("stoc_minim"))))
            self.supplier_terms(values[2])
            exists = self.conn.execute("SELECT 1 FROM products WHERE code = ?", (code,)).fetchone()
            self.conn.execute(
                "INSERT INTO products(code, name, category, supplier, sale_price, cost, min_stock) "
                "VALUES (?, ?, ?, ?, ?, ?, ?) ON CONFLICT(code) DO UPDATE SET name = excluded.name, "
                "category = excluded.category, supplier = excluded.supplier, "
                "sale_price = excluded.sale_price, cost = excluded.cost, min_stock = excluded.min_stock",
                (code, *values))
            if not exists:
                self.add("products")

    def stock(self, table: Table, filename: str) -> None:
        """Stocul de deschidere: din foaia Stocuri (prima lună) sau din inventarul CSV."""
        initial_col = table.col("stoc_initial")
        entries_col = table.col("intrari")
        day = month_key = None
        match = re.search(r"stoc_initial_(\d{2})_(\d{2})", initial_col or "")
        if match:  # inventar CSV: „stoc_initial_01_09”, anul se ia din numele fișierului
            year = re.search(r"(20\d{2})", filename)
            year = int(year.group(1)) if year else date.today().year
            day = f"{year}-{match.group(2)}-{match.group(1)}"
        seen_codes = set()
        for row in table.rows:
            code = str(row.get("cod_produs") or "").strip()
            if not code or code in seen_codes:
                continue
            opening = row.get(initial_col)
            if not is_number(opening):  # în foaia Stocuri, lunile următoare sunt formule
                continue
            seen_codes.add(code)
            row_day = day
            if row.get("luna"):
                month_key = str(row["luna"])[:7]
                row_day = f"{month_key}-01"
            cost = to_float(row.get(table.col("cost_unitar")))
            pid = self.product_id(str(row.get("produs") or code), code, str(row.get("categorie") or ""), cost)
            if cost:
                self.conn.execute("UPDATE products SET cost = ? WHERE id = ? AND cost = 0", (cost, pid))
            # Un stoc de deschidere mai vechi (cu istoricul de după el) nu se înlocuiește.
            current = self.conn.execute("SELECT opening_date FROM products WHERE id = ?", (pid,)).fetchone()
            if current["opening_date"] and row_day and current["opening_date"] <= row_day:
                continue
            self.conn.execute("UPDATE products SET opening_stock = ?, opening_date = ? WHERE id = ?",
                              (int(to_float(opening)), row_day, pid))
            self.add("stock")
            # Inventarul CSV are doar totalul intrărilor: devine o recepție sumară la final de lună.
            if entries_col and is_number(row.get(entries_col)) and row_day:
                start, end = month_bounds(row_day[:7])
                quantity = int(to_float(row.get(entries_col)))
                has_nir = self.conn.execute(
                    "SELECT 1 FROM purchases WHERE product_id = ? AND date BETWEEN ? AND ?",
                    (pid, start, end)).fetchone()
                if quantity and not has_nir:
                    self.conn.execute(
                        "INSERT INTO purchases(date, doc_number, supplier, product_id, quantity, unit_cost, "
                        "due_date, paid) VALUES (?, ?, ?, ?, ?, ?, ?, 1)",
                        (end, f"INV-{row_day[:7]}", "Intrări din inventar", pid, quantity, cost, end))
                    self.add("purchases")

    def purchases(self, table: Table) -> None:
        dates = [to_date(row.get("data")) for row in table.rows if row.get("data")]
        reference = self.reference or (month_bounds(max(dates)[:7])[1] if dates else None)
        for row in table.rows:
            code = str(row.get("cod_produs") or "").strip()
            day = to_date(row.get("data"))
            if not code or not day:
                continue
            cost = to_float(row.get(table.col("cost_unitar")))
            pid = self.product_id(str(row.get("produs") or code), code, cost=cost)
            supplier = str(row.get("furnizor") or "")
            doc = str(row.get("nr_nir") or "")
            if supplier:
                self.conn.execute("UPDATE products SET supplier = ? WHERE id = ? AND supplier = ''",
                                  (supplier, pid))
            if self.conn.execute("SELECT 1 FROM purchases WHERE doc_number = ? AND product_id = ?",
                                 (doc, pid)).fetchone():
                continue
            due = (date.fromisoformat(day) + timedelta(days=self.supplier_terms(supplier))).isoformat()
            # Ca la clienți: ce are scadența după data de referință e încă de plătit.
            paid = 1 if reference is None or due <= reference else 0
            self.conn.execute(
                "INSERT INTO purchases(date, doc_number, supplier, product_id, quantity, unit_cost, "
                "due_date, paid) VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
                (day, doc, supplier, pid, int(to_float(row.get("cantitate"))), cost, due, paid))
            self.add("purchases")

    def sales(self, table: Table) -> None:
        dates = [to_date(row.get("data")) for row in table.rows if row.get("data")]
        reference = self.reference or (month_bounds(max(dates)[:7])[1] if dates else None)
        price_col, cost_col = table.col("pret_unitar"), table.col("cost_achizitie")
        for row in table.rows:
            invoice = str(row.get("nr_factura") or "").strip()
            day = to_date(row.get("data"))
            if not day:
                continue
            if invoice and self.conn.execute("SELECT 1 FROM sales WHERE invoice_number = ?",
                                             (invoice,)).fetchone():
                continue
            name = str(row.get("produs") or "").strip()
            quantity = int(to_float(row.get("cantitate")) or 1)
            price = to_float(row.get(price_col))
            cost = to_float(row.get(cost_col))
            pid = self.product_id(name, category=str(row.get("categorie") or ""), cost=cost, price=price)
            self.conn.execute("UPDATE products SET sale_price = ? WHERE id = ? AND sale_price = 0", (price, pid))
            client = self.client(str(row.get("client") or ""))
            due = (date.fromisoformat(day) + timedelta(days=client["payment_terms"])).isoformat()
            # Ca în raportul contabil: o factură cu scadența după data de referință e neîncasată.
            paid = 1 if reference is None or due <= reference else 0
            self.conn.execute(
                "INSERT INTO sales(date, description, product, quantity, amount, invoice_number, product_id, "
                "client_id, unit_price, unit_cost, due_date, paid) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                (day, f"{name} ×{quantity}", name, quantity, round(quantity * price, 2), invoice or None,
                 pid, client["id"], price, cost, due, paid))
            self.add("sales")

    def expenses(self, table: Table) -> None:
        vat_col = table.col("cota_tva")
        amount_col = table.col("suma_fara_tva")
        defaults = {row["name"]: row["vat_kind"] for row in self.conn.execute("SELECT * FROM categories")}
        for row in table.rows:
            day = to_date(row.get("data"))
            if not day or not str(row.get("categorie") or "").strip():
                continue
            category = self.category(row["categorie"])
            description = str(row.get("descriere") or "")
            amount = round(to_float(row.get(amount_col)), 2)
            raw_vat = row.get(vat_col) if vat_col else None
            if isinstance(raw_vat, str) and "$B$5" in raw_vat:
                vat_kind = "reduced"
            elif isinstance(raw_vat, str) and "$B$4" in raw_vat:
                vat_kind = "standard"
            elif vat_col and (raw_vat in (None, "", "=0") or (is_number(raw_vat) and to_float(raw_vat) == 0)):
                vat_kind = "none"
            elif is_number(raw_vat) and to_float(raw_vat) < 0.15:
                vat_kind = "reduced"
            elif "apa" in norm(description).split("_"):
                vat_kind = "reduced"
            else:
                vat_kind = defaults.get(category, "standard")
            if self.conn.execute(
                    "SELECT 1 FROM expenses WHERE date = ? AND category = ? AND description = ? AND amount = ?",
                    (day, category, description, amount)).fetchone():
                continue
            self.conn.execute(
                "INSERT INTO expenses(date, category, description, amount, supplier, vat_kind) "
                "VALUES (?, ?, ?, ?, ?, ?)",
                (day, category, description, amount, str(row.get("furnizor") or ""), vat_kind))
            self.add("expenses")

    def payroll(self, table: Table) -> None:
        gross_col = table.col("salariu_brut")
        first_data = self.conn.execute("SELECT MIN(date) FROM sales").fetchone()[0]
        default_start = f"{(first_data or date.today().isoformat())[:7]}-01"
        for row in table.rows:
            name = str(row.get("nume") or "").strip()
            if not name:
                continue
            gross = to_float(row.get(gross_col))
            start = f"{str(row['luna'])[:7]}-01" if row.get("luna") else default_start
            employee = self.conn.execute("SELECT * FROM employees WHERE name = ?", (name,)).fetchone()
            if not employee:
                cursor = self.conn.execute(
                    "INSERT INTO employees(name, role, hire_date) VALUES (?, ?, ?)",
                    (name, str(row.get("functie") or ""), start))
                self.conn.execute(
                    "INSERT INTO salary_history(employee_id, gross, effective_from, note) "
                    "VALUES (?, ?, ?, 'Salariu inițial')", (cursor.lastrowid, gross, start))
                self.add("employees")
                continue
            if start < employee["hire_date"]:
                self.conn.execute("UPDATE employees SET hire_date = ? WHERE id = ?", (start, employee["id"]))
                self.conn.execute("UPDATE salary_history SET effective_from = ? WHERE employee_id = ? "
                                  "AND effective_from = ?", (start, employee["id"], employee["hire_date"]))
            last = self.conn.execute(
                "SELECT gross FROM salary_history WHERE employee_id = ? AND effective_from <= ? "
                "ORDER BY effective_from DESC LIMIT 1", (employee["id"], start)).fetchone()
            if last and last["gross"] != gross:
                self.conn.execute(
                    "INSERT INTO salary_history(employee_id, gross, effective_from, note) "
                    "VALUES (?, ?, ?, 'Modificare salariu')", (employee["id"], gross, start))

    def run(self, tables: list[tuple[Table, str]]) -> dict:
        recognized = [(table, filename, table.kind()) for table, filename in tables]
        for table, _, kind in recognized:
            if kind is None:
                self.warnings.append(f"Nu am recunoscut tabelul „{table.name}”; l-am sărit.")
        for step in self.ORDER:
            for table, filename, kind in recognized:
                if kind == step:
                    if kind == "stock":
                        self.stock(table, filename)
                    else:
                        getattr(self, kind)(table)
        return {"counts": self.counts, "warnings": self.warnings}


def clear_business_data(conn: sqlite3.Connection) -> None:
    for table in DATA_TABLES:
        conn.execute(f"DELETE FROM {table}")


def import_files(conn: sqlite3.Connection, files: list[tuple[str, bytes]], replace: bool = False) -> dict:
    """Importă o listă de fișiere (nume, conținut). Totul sau nimic: la eroare nu se salvează nimic."""
    tables, reference = [], None
    for filename, content in files:
        lower = filename.lower()
        if lower.endswith(".xlsx"):
            found, ref = tables_from_xlsx(filename, content)
            reference = reference or ref
            tables += [(table, filename) for table in found]
        elif lower.endswith(".csv"):
            tables += [(table, filename) for table in tables_from_csv(filename, content)]
        else:
            raise ValueError(f"„{filename}” nu este CSV sau XLSX.")
    if not tables:
        raise ValueError("Nu am găsit niciun tabel de importat.")
    conn.execute("SAVEPOINT import_data")
    try:
        if replace:
            clear_business_data(conn)
        result = Importer(conn, reference).run(tables)
        conn.execute("RELEASE SAVEPOINT import_data")
        return result
    except Exception:
        conn.execute("ROLLBACK TO SAVEPOINT import_data")
        conn.execute("RELEASE SAVEPOINT import_data")
        raise
